import base64
import io
import json
import os
import unittest
import zipfile
from unittest.mock import patch
from PIL import Image, ImageDraw
from pptx import Presentation
from marketplace import ai_copy, baseline, seller_assets, validate_copy, theme_from_reference

class MarketplaceTests(unittest.TestCase):
    def setUp(self):
        stream=io.BytesIO();Image.new('RGB',(600,900),'#336633').save(stream,'PNG');self.art=stream.getvalue()
    def test_basic_seller_assets_are_editable_and_full_sized(self):
        result=seller_assets(self.art,{'title':'Forest','shop':'Bookish Daisy'},mode='basic')
        self.assertEqual(result['source'],'Basic template — not AI');self.assertEqual(len(result['copy']['tags']),13)
        with zipfile.ZipFile(io.BytesIO(base64.b64decode(result['seller_zip']))) as z:
            self.assertIn('Shop-copy.txt',z.namelist());self.assertIn('Canva-editable.pptx',z.namelist());self.assertNotIn('Original-art.png',z.namelist())
        presentation=Presentation(io.BytesIO(base64.b64decode(result['canva_pptx'])))
        self.assertEqual(len(presentation.slides),3)
        self.assertTrue(any(shape.has_text_frame and 'Forest' in shape.text for shape in presentation.slides[0].shapes))
        self.assertTrue(any(shape.shape_type==13 for shape in presentation.slides[0].shapes))
        for img in result['listing_images']:
            with Image.open(io.BytesIO(base64.b64decode(img['data']))) as image:self.assertEqual(image.size,(2000,2000))
    def test_missing_ai_does_not_invent_copy_or_block_images(self):
        with patch.dict(os.environ,{},clear=True):result=seller_assets(self.art,{'title':'Forest'},mode='ai')
        self.assertIsNone(result['copy']);self.assertIn('not connected',result['copy_error']);self.assertEqual(len(result['listing_images']),3)
        with zipfile.ZipFile(io.BytesIO(base64.b64decode(result['seller_zip']))) as z:
            self.assertIn('AI-status.txt',z.namelist());self.assertNotIn('Shop-copy.txt',z.namelist())
    def test_gemini_receives_image_and_returns_structured_copy(self):
        copy=baseline('Forest')
        reply={'candidates':[{'finishReason':'STOP','content':{'parts':[{'text':json.dumps(copy)}]}}]}
        class Response(io.BytesIO):
            def __enter__(self):return self
            def __exit__(self,*args):self.close()
        def respond(request,timeout):
            self.assertEqual(request.full_url,'https://generativelanguage.googleapis.com/v1beta/models/gemini-2.5-flash:generateContent')
            body=json.loads(request.data);self.assertEqual(body['contents'][0]['parts'][1]['inlineData']['mimeType'],'image/png')
            self.assertEqual(timeout,75);return Response(json.dumps(reply).encode())
        with patch('marketplace.urllib.request.urlopen',side_effect=respond):
            result=ai_copy(Image.open(io.BytesIO(self.art)),'Forest',{},supplied_key='test-only-key')
        self.assertEqual(result['title'],copy['title'])
    def test_banner_zip_theme_sampling_and_ambiguity(self):
        image=Image.new('RGB',(1200,300),'#f7e7da');draw=ImageDraw.Draw(image);draw.rectangle((20,20,300,260),fill='#703a56');draw.rectangle((320,20,600,260),fill='#24212a')
        data=io.BytesIO();image.save(data,'PNG');archive=io.BytesIO()
        with zipfile.ZipFile(archive,'w') as z:z.writestr('brand/shop-banner.png',data.getvalue())
        result=theme_from_reference(archive.getvalue());self.assertEqual(result['source'],'brand/shop-banner.png');self.assertEqual(set(result['palette']),{'background','accent','text'})
        for value in result['palette'].values():self.assertRegex(value,r'^#[0-9a-f]{6}$')
        with Image.open(io.BytesIO(base64.b64decode(result['banner']))) as banner:self.assertEqual(banner.size,(1200,300))
        archive=io.BytesIO()
        with zipfile.ZipFile(archive,'w') as z:
            z.writestr('banner1.png',data.getvalue());z.writestr('banner2.png',data.getvalue())
        with self.assertRaisesRegex(ValueError,'Several banner'):theme_from_reference(archive.getvalue())
    def test_malformed_tags_are_rejected(self):
        for tags in [['one'],['duplicate']*13,['a'*21]*13]:
            value=baseline('Forest');value['tags']=tags
            with self.assertRaises(ValueError):validate_copy(value)

if __name__=='__main__':unittest.main()
